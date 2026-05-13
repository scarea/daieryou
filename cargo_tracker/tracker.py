#!/usr/bin/env python3
"""
Ethiopian Cargo AWB tracker.

Reads AWB numbers from an Excel file, submits the public website form one by
one, and writes flight schedule / shipment history rows to a result workbook.
"""

from __future__ import annotations

import argparse
import json
import random
import re
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


TRACKING_URL = "https://ethiopiancargo.azurewebsites.net/my-cargo/track-your-shipment/Index/"
DEFAULT_MIN_DELAY_SECONDS = 6
DEFAULT_MAX_DELAY_SECONDS = 18
DEFAULT_RETRIES = 3


@dataclass
class FlightRow:
    awb: str
    route: str
    flight_no: str
    flight_date: str
    cargo_status: str


@dataclass
class HistoryRow:
    awb: str
    event: str
    detail: str
    event_time: str


@dataclass
class TrackResult:
    awb: str
    status: str
    flights: list[FlightRow]
    histories: list[HistoryRow]
    error: str = ""


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def normalize_awb(value: object) -> str:
    text = str(value or "").strip()
    text = text.replace(" ", "")
    return text


def read_awbs(input_path: Path) -> list[str]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    sheet = workbook.active

    awbs: list[str] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        awb = normalize_awb(row[0] if row else "")
        if awb:
            awbs.append(awb)

    # Deduplicate while preserving order, so the same AWB is not queried twice.
    seen: set[str] = set()
    unique_awbs: list[str] = []
    for awb in awbs:
        if awb not in seen:
            seen.add(awb)
            unique_awbs.append(awb)

    return unique_awbs


def parse_tracking_html(awb: str, html: str) -> TrackResult:
    soup = BeautifulSoup(html, "html.parser")

    page_text = normalize_text(soup.get_text(" "))
    if "Flight Schedule" not in page_text and "Shipment History" not in page_text:
        return TrackResult(awb=awb, status="无结果", flights=[], histories=[], error="页面未返回航班或轨迹信息")

    flights: list[FlightRow] = []
    for card in soup.select(".flight-card"):
        route = normalize_text(card.select_one(".route-text").get_text(" ") if card.select_one(".route-text") else "")
        meta = normalize_text(card.select_one(".flight-meta").get_text(" ") if card.select_one(".flight-meta") else "")
        cargo_status = normalize_text(card.select_one(".status-badge").get_text(" ") if card.select_one(".status-badge") else "")

        flight_no = ""
        flight_date = ""
        if "•" in meta:
            parts = [normalize_text(part) for part in meta.split("•", 1)]
            flight_no = parts[0]
            flight_date = parts[1] if len(parts) > 1 else ""
        else:
            flight_no = meta

        if route or flight_no or flight_date or cargo_status:
            flights.append(
                FlightRow(
                    awb=awb,
                    route=route,
                    flight_no=flight_no,
                    flight_date=flight_date,
                    cargo_status=cargo_status,
                )
            )

    histories: list[HistoryRow] = []
    for container in soup.select(".status-container"):
        event_el = container.select_one(".status-font")
        event = normalize_text(event_el.get_text(" ") if event_el else "")

        time_col = container.select_one(".col-sm-12.col-md-6.col-lg-4")
        strongs = [normalize_text(item.get_text(" ")) for item in (time_col.select("strong") if time_col else container.select("strong"))]
        event_time = ""
        for text in reversed(strongs):
            if re.search(r"\d{2}-[A-Za-z]{3}-\d{2}", text):
                event_time = text
                break

        detail = ""
        detail_col = container.select_one(".col-sm-12.col-md-6.col-lg-8")
        detail_candidates = detail_col.find_all("div", recursive=False) if detail_col else container.select("div")
        for div in detail_candidates:
            text = normalize_text(div.get_text(" "))
            if "Pcs" in text and "from" in text:
                detail = text
                break

        if event or detail or event_time:
            histories.append(HistoryRow(awb=awb, event=event, detail=detail, event_time=event_time))

    if not flights and not histories:
        return TrackResult(awb=awb, status="无结果", flights=[], histories=[], error="页面结构变化或没有可解析信息")

    return TrackResult(awb=awb, status="成功", flights=flights, histories=histories)


def should_stop_for_protection(response: requests.Response, html: str) -> bool:
    if response.status_code in {403, 429}:
        return True

    lowered = html.lower()
    protection_markers = [
        "too many requests",
        "access denied",
        "request blocked",
        "rate limit",
        "verify you are human",
    ]
    return any(marker in lowered for marker in protection_markers)


def query_one_awb(
    session: requests.Session,
    awb: str,
    retries: int,
    min_delay: float,
    max_delay: float,
) -> TrackResult:
    last_error = ""

    for attempt in range(1, retries + 1):
        try:
            if attempt == 1:
                session.get(TRACKING_URL, timeout=30)

            response = session.post(
                TRACKING_URL,
                data={"AirwayBilNum": awb},
                headers={
                    "Content-Type": "application/x-www-form-urlencoded",
                    "Referer": TRACKING_URL,
                },
                timeout=45,
            )

            html = response.text
            if should_stop_for_protection(response, html):
                raise RuntimeError(f"网站返回保护/限流信号，HTTP {response.status_code}，本次任务已停止")

            response.raise_for_status()
            return parse_tracking_html(awb, html)
        except RuntimeError:
            raise
        except Exception as exc:
            last_error = str(exc)
            if attempt < retries:
                backoff = (15 * attempt) + random.uniform(min_delay, max_delay)
                print(f"[{awb}] 第 {attempt} 次失败，{round(backoff, 1)} 秒后重试：{last_error}")
                time.sleep(backoff)

    return TrackResult(awb=awb, status="失败", flights=[], histories=[], error=last_error)


def autosize_columns(sheet) -> None:
    for column_cells in sheet.columns:
        max_len = 0
        column = column_cells[0].column
        for cell in column_cells:
            max_len = max(max_len, len(str(cell.value or "")))
        sheet.column_dimensions[get_column_letter(column)].width = min(max(max_len + 2, 12), 60)


def write_results(output_path: Path, results: Iterable[TrackResult]) -> None:
    workbook = Workbook()

    summary = workbook.active
    summary.title = "查询汇总"
    flight_sheet = workbook.create_sheet("航班计划")
    history_sheet = workbook.create_sheet("货物轨迹")

    summary_headers = ["AWB", "查询状态", "航班计划数量", "轨迹数量", "错误信息"]
    flight_headers = ["AWB", "航段", "航班号", "日期", "货运状态"]
    history_headers = ["AWB", "事件", "详情", "时间"]

    summary.append(summary_headers)
    flight_sheet.append(flight_headers)
    history_sheet.append(history_headers)

    header_fill = PatternFill("solid", fgColor="1F7A3A")
    header_font = Font(color="FFFFFF", bold=True)

    for sheet in [summary, flight_sheet, history_sheet]:
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        sheet.freeze_panes = "A2"

    for result in results:
        summary.append([result.awb, result.status, len(result.flights), len(result.histories), result.error])

        for flight in result.flights:
            flight_sheet.append([flight.awb, flight.route, flight.flight_no, flight.flight_date, flight.cargo_status])

        for history in result.histories:
            history_sheet.append([history.awb, history.event, history.detail, history.event_time])

    for sheet in [summary, flight_sheet, history_sheet]:
        sheet.auto_filter.ref = sheet.dimensions
        autosize_columns(sheet)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def create_template(output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "AWB列表"
    sheet.append(["AWB"])
    sheet.append(["071-61046882"])
    sheet["A1"].fill = PatternFill("solid", fgColor="1F7A3A")
    sheet["A1"].font = Font(color="FFFFFF", bold=True)
    sheet.column_dimensions["A"].width = 24
    workbook.save(output_path)


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Ethiopian Cargo AWB 批量查询工具")
    parser.add_argument("--input", "-i", type=Path, help="输入 Excel 文件，第一列从第二行开始放 AWB")
    parser.add_argument("--output", "-o", type=Path, help="输出 Excel 文件")
    parser.add_argument("--create-template", type=Path, help="创建输入模板 Excel")
    parser.add_argument("--min-delay", type=float, default=DEFAULT_MIN_DELAY_SECONDS, help="每票查询后的最小等待秒数")
    parser.add_argument("--max-delay", type=float, default=DEFAULT_MAX_DELAY_SECONDS, help="每票查询后的最大等待秒数")
    parser.add_argument("--retries", type=int, default=DEFAULT_RETRIES, help="单票失败后的最大尝试次数")
    parser.add_argument("--fast-test", action="store_true", help="仅测试时使用：缩短随机等待")
    return parser


def main() -> int:
    args = build_arg_parser().parse_args()

    if args.create_template:
        create_template(args.create_template)
        print(f"已创建模板：{args.create_template}")
        return 0

    if not args.input:
        print("请提供 --input 输入 Excel 文件。")
        return 2

    input_path = args.input
    if not input_path.exists():
        print(f"输入文件不存在：{input_path}")
        return 2

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = args.output or Path("output") / f"ethiopian_cargo_result_{timestamp}.xlsx"

    min_delay = 1 if args.fast_test else args.min_delay
    max_delay = 2 if args.fast_test else args.max_delay

    if min_delay > max_delay:
        print("--min-delay 不能大于 --max-delay")
        return 2

    awbs = read_awbs(input_path)
    if not awbs:
        print("输入 Excel 中没有找到 AWB。请确认第一列从第二行开始填写。")
        return 2

    print(f"读取到 {len(awbs)} 个不重复 AWB。")

    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": "CargoTrackingBackOfficeTool/1.0",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        }
    )

    results: list[TrackResult] = []
    try:
        for index, awb in enumerate(awbs, start=1):
            print(f"[{index}/{len(awbs)}] 查询 {awb} ...")
            result = query_one_awb(session, awb, args.retries, min_delay, max_delay)
            results.append(result)
            write_results(output_path, results)
            print(f"[{awb}] {result.status}，航班 {len(result.flights)} 条，轨迹 {len(result.histories)} 条。")

            if index < len(awbs):
                delay = random.uniform(min_delay, max_delay)
                print(f"等待 {round(delay, 1)} 秒后继续下一票。")
                time.sleep(delay)
    except RuntimeError as exc:
        if results:
            write_results(output_path, results)
        print(f"已停止：{exc}")
        print(f"已保存当前进度：{output_path}")
        return 1

    metadata = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "input": str(input_path),
        "output": str(output_path),
        "count": len(awbs),
    }
    output_path.with_suffix(".json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"完成，结果文件：{output_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
